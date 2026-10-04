"""Excel出力(日本語見出し・列幅調整)のテスト。"""

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from audit_rules import audit_records, read_csv  # noqa: E402
from excel_report import (  # noqa: E402
    JAPANESE_LABELS,
    MAX_COLUMN_WIDTH,
    calculate_column_width,
    display_value,
    display_width,
    write_audit_xlsx,
)


class ExcelReportTest(unittest.TestCase):
    """英語の項目名の下に日本語見出しが付き、列幅が調整されることを確認する。"""

    @classmethod
    def setUpClass(cls) -> None:
        records = audit_records(read_csv(PROJECT_ROOT / "data" / "mock_ai_output.csv"))
        cls.records = records
        cls.columns = list(records[0].keys())
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "audit_result.xlsx"
        write_audit_xlsx(cls.path, records)
        cls.sheet = load_workbook(cls.path).active

    @classmethod
    def tearDownClass(cls) -> None:
        cls.tmp.cleanup()

    def test_every_column_has_a_japanese_label(self) -> None:
        for column in self.columns:
            self.assertIn(column, JAPANESE_LABELS)

    def test_row1_english_and_row2_japanese(self) -> None:
        for index, column in enumerate(self.columns, start=1):
            self.assertEqual(self.sheet.cell(row=1, column=index).value, column)
            self.assertEqual(
                self.sheet.cell(row=2, column=index).value, JAPANESE_LABELS[column]
            )

    def test_key_headings_are_the_agreed_japanese_words(self) -> None:
        self.assertEqual(JAPANESE_LABELS["record_id"], "識別番号")
        self.assertEqual(JAPANESE_LABELS["extracted_assignee"], "抽出した担当者")
        self.assertEqual(JAPANESE_LABELS["audit_status"], "項目の充足")
        self.assertEqual(JAPANESE_LABELS["overall_status"], "総合判定")

    def test_all_records_written_from_row3(self) -> None:
        self.assertEqual(self.sheet.max_row, 2 + len(self.records))
        self.assertEqual(self.sheet.cell(row=3, column=1).value, self.records[0]["record_id"])

    def test_column_widths_fit_content_with_cap(self) -> None:
        for index in range(1, len(self.columns) + 1):
            letter = self.sheet.cell(row=1, column=index).column_letter
            width = self.sheet.column_dimensions[letter].width
            self.assertGreaterEqual(width, 8)
            self.assertLessEqual(width, MAX_COLUMN_WIDTH)
        # 申し送り原文(長い文章)は、2行に折り返して全文が入る幅まで縮める
        self.assertLessEqual(self._width("message_text"), 46)

    def _width(self, column: str) -> float:
        index = self.columns.index(column) + 1
        letter = self.sheet.cell(row=1, column=index).column_letter
        return self.sheet.column_dimensions[letter].width

    def test_short_content_columns_are_not_too_wide(self) -> None:
        # 内容が短い列(部署・担当者・日時など)は、見出しのせいで広がりすぎない
        self.assertLessEqual(self._width("source_department"), 13)
        self.assertLessEqual(self._width("extracted_assignee"), 17)
        self.assertLessEqual(self._width("extracted_priority"), 14)
        self.assertLessEqual(self._width("record_id"), 11)

    def test_headings_are_not_cut_off(self) -> None:
        # 日本語見出しは1行で入る幅、英語は8ptで入る幅を必ず確保する
        for column in self.columns:
            japanese = JAPANESE_LABELS[column]
            self.assertGreaterEqual(self._width(column), display_width(japanese) + 2)
            self.assertGreaterEqual(self._width(column), display_width(column) * 0.7)

    def test_sample_rows_fit_within_two_lines(self) -> None:
        # 試験データは、どの列も2行以内に収まる(3行以上になると行が高くなり余白が目立つ)
        for column in self.columns:
            longest = max(
                display_width(display_value(column, record.get(column, "")))
                for record in self.records
            )
            self.assertLessEqual(longest, 2 * (self._width(column) - 1), column)

    def test_data_row_heights_are_uniform_for_sample(self) -> None:
        heights = {
            self.sheet.row_dimensions[r].height for r in range(3, self.sheet.max_row + 1)
        }
        self.assertEqual(heights, {40})

    def test_long_text_gets_taller_row(self) -> None:
        long_record = dict(self.records[0], message_text="あ" * 400)
        path = Path(self.tmp.name) / "long.xlsx"
        write_audit_xlsx(path, [long_record])
        sheet = load_workbook(path).active
        self.assertGreater(sheet.row_dimensions[3].height, 40)

    def test_short_data_gets_single_line_rows(self) -> None:
        short_record = {"record_id": "H1", "audit_status": "READY"}
        path = Path(self.tmp.name) / "short.xlsx"
        write_audit_xlsx(path, [short_record])
        sheet = load_workbook(path).active
        self.assertEqual(sheet.row_dimensions[3].height, 20)

    def test_cells_are_vertically_centered(self) -> None:
        self.assertEqual(self.sheet.cell(row=3, column=1).alignment.vertical, "center")

    def test_freeze_panes_and_filter(self) -> None:
        self.assertEqual(self.sheet.freeze_panes, "B3")
        self.assertTrue(self.sheet.auto_filter.ref.startswith("A2:"))

    def test_status_values_show_english_and_japanese(self) -> None:
        self.assertEqual(display_value("overall_status", "NEEDS_REVIEW"), "NEEDS_REVIEW(要確認)")
        self.assertEqual(display_value("overall_status", "READY"), "READY(問題なし)")
        self.assertEqual(display_value("overall_status", "INFO_ONLY"), "INFO_ONLY(共有のみ)")
        # 「項目の充足」は、総合判定と間違えないよう別の言葉で表示する
        self.assertEqual(display_value("audit_status", "NEEDS_REVIEW"), "NEEDS_REVIEW(不足あり)")
        self.assertEqual(display_value("audit_status", "READY"), "READY(そろっている)")
        self.assertEqual(display_value("sla_audit_status", "ON_TIME"), "ON_TIME(期限内)")
        self.assertEqual(display_value("action_required", "true"), "true(対応が必要)")

    def test_unmapped_values_pass_through(self) -> None:
        self.assertEqual(display_value("extracted_priority", "高"), "高")
        self.assertEqual(display_value("message_text", "READY"), "READY")
        self.assertEqual(display_value("audit_status", "OTHER"), "OTHER")

    def test_sheet_cells_show_both_languages(self) -> None:
        col = self.columns.index("overall_status") + 1
        shown = {self.sheet.cell(row=r, column=col).value for r in range(3, self.sheet.max_row + 1)}
        self.assertTrue(all("(" in v for v in shown))
        self.assertIn("NEEDS_REVIEW(要確認)", shown)
        fill_col = self.columns.index("audit_status") + 1
        fill_shown = {
            self.sheet.cell(row=r, column=fill_col).value for r in range(3, self.sheet.max_row + 1)
        }
        self.assertTrue(all("(" in v for v in fill_shown))
        self.assertNotIn("READY(問題なし)", fill_shown)

    def test_status_column_wide_enough_for_both_languages(self) -> None:
        col = self.columns.index("overall_status") + 1
        letter = self.sheet.cell(row=1, column=col).column_letter
        self.assertGreaterEqual(
            self.sheet.column_dimensions[letter].width, display_width("NEEDS_REVIEW(要確認)")
        )

    def test_display_width_counts_fullwidth_as_two(self) -> None:
        self.assertEqual(display_width("abc"), 3)
        self.assertEqual(display_width("担当者"), 6)

    def test_short_column_gets_minimum_width(self) -> None:
        self.assertEqual(calculate_column_width("a", "あ", ["b"]), 8.0)

    def test_empty_records_rejected(self) -> None:
        with self.assertRaises(ValueError):
            write_audit_xlsx(Path(self.tmp.name) / "x.xlsx", [])


class ExcelFormulaSafetyTest(unittest.TestCase):
    """外部由来の文字列(申し送りの原文など)が、Excelの数式として保存されないことを確認する。"""

    DANGEROUS = ["=1+1", "=HYPERLINK(\"http://example.com\",\"x\")", "+1+1", "-1+1", "@SUM(1,1)"]

    def _write(self, text: str) -> Path:
        record = audit_records(read_csv(PROJECT_ROOT / "data" / "mock_ai_output.csv"))[0]
        record = dict(record, message_text=text, extracted_action=text)
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "out.xlsx"
        write_audit_xlsx(path, [record])
        return path

    def test_text_starting_with_formula_characters_is_stored_as_text(self) -> None:
        for text in self.DANGEROUS:
            with self.subTest(text=text):
                path = self._write(text)
                sheet = load_workbook(path).active
                values = [cell for row in sheet.iter_rows() for cell in row if cell.value == text]
                self.assertTrue(values, "文字として残っていること")
                for cell in values:
                    self.assertEqual(cell.data_type, "s")

    def test_no_formula_element_in_saved_xml(self) -> None:
        path = self._write("=1+1")
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertNotIn("<f>", xml)


if __name__ == "__main__":
    unittest.main()
