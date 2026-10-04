"""監査結果をExcelで読みやすい形(日本語見出し・列幅調整つき)へ出力する。"""

from __future__ import annotations

import csv
import math
import unicodedata
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# 1. 見出しの日本語表記
# ============================================================

# CSVの項目名(英語)と、Excelに表示する日本語見出しの対応表。
# CSV自体の項目名は変えない(評価・テストが英語名で読むため)。
JAPANESE_LABELS: Dict[str, str] = {
    "record_id": "識別番号",
    "submitted_at": "登録日時",
    "source_department": "発信部署",
    "message_text": "申し送り原文",
    "extracted_action": "抽出した作業内容",
    "extracted_assignee": "抽出した担当者",
    "extracted_deadline": "抽出した期限",
    "extracted_priority": "AIの緊急度",
    "action_required": "対応要否",
    "rule_priority": "ルール判定の緊急度",
    "priority_audit_status": "緊急度の照合結果",
    "priority_audit_message": "緊急度の照合理由",
    "sla_audit_status": "期限超過の判定",
    "sla_audit_message": "期限超過の理由",
    "duplicate_audit_status": "重複の判定",
    "duplicate_audit_message": "重複の理由",
    "vague_audit_status": "作業内容の具体性",
    "vague_audit_message": "作業内容の具体性の理由",
    "audit_status": "項目の充足",
    "missing_fields": "不足項目",
    "audit_message": "項目の充足の理由",
    "overall_status": "総合判定",
    "overall_message": "総合判定の理由",
}

# 判定結果などの値は「英語(日本語)」の形で表示する。
# 英語の値はCSVと同じなので、英語でも日本語でも探せる。対応表にない値はそのまま表示する。
_REVIEW_STATUS = {
    "READY": "問題なし",
    "NEEDS_REVIEW": "要確認",
    "INFO_ONLY": "共有のみ",
}
# 「項目の充足」は、作業内容・担当者・期限がそろっているかだけを表す(総合判定とは別)
_FILL_STATUS = {
    "READY": "そろっている",
    "NEEDS_REVIEW": "不足あり",
    "INFO_ONLY": "共有のみ",
}
_AUDIT_AUX_STATUS = {
    "MATCH": "一致",
    "NEEDS_REVIEW": "要確認",
    "ON_TIME": "期限内",
    "UNIQUE": "重複なし",
    "NOT_APPLICABLE": "判定対象外",
    "UNKNOWN": "判定不能",
    "CLEAR": "具体的",
}
VALUE_LABELS: Dict[str, Dict[str, str]] = {
    "audit_status": _FILL_STATUS,
    "overall_status": _REVIEW_STATUS,
    "priority_audit_status": _AUDIT_AUX_STATUS,
    "sla_audit_status": _AUDIT_AUX_STATUS,
    "duplicate_audit_status": _AUDIT_AUX_STATUS,
    "vague_audit_status": _AUDIT_AUX_STATUS,
    "action_required": {"true": "対応が必要", "false": "対応不要"},
}

MIN_COLUMN_WIDTH = 8
ENGLISH_WIDTH_RATIO = 0.7
ROW_HEIGHT_PER_LINE = 20
MAX_COLUMN_WIDTH = 50
# 内容がこの幅を超える列は、2行に折り返して列幅を縮める
SINGLE_LINE_LIMIT = 40
TARGET_LINES = 2
ENGLISH_ROW = 1
JAPANESE_ROW = 2
FIRST_DATA_ROW = 3


def display_value(column: str, value: object) -> str:
    """値を「英語(日本語)」の表示へ変換する。対応表にない値はそのまま返す。"""

    text = str(value)
    japanese = VALUE_LABELS.get(column, {}).get(text)
    return f"{text}({japanese})" if japanese else text


# ============================================================
# 2. 列幅の計算
# ============================================================

def display_width(text: str) -> int:
    """全角を2、半角を1として表示幅を数える。"""

    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in text)


def calculate_column_width(english: str, japanese: str, values: List[str]) -> float:
    """内容を基準に、見出しが読める最小限の列幅(上限あり)を返す。

    日本語見出しは(太字なので少し余裕を足して)1行に収める。英語の項目名は小さい
    文字(8pt)で表示するため、実際の文字数の約7割の幅で足りる。
    """

    longest = max([display_width(text) for text in values] or [0])
    if longest + 2 > SINGLE_LINE_LIMIT:
        # 長い文は、2行に折り返せる幅(半分+余裕)にして、列が横に伸びすぎないようにする
        content = math.ceil(longest / TARGET_LINES) + 3
    else:
        content = longest + 2
    japanese_need = display_width(japanese) + 3
    english_need = math.ceil(display_width(english) * ENGLISH_WIDTH_RATIO) + 1
    width = max(content, japanese_need, english_need, MIN_COLUMN_WIDTH)
    return float(min(width, MAX_COLUMN_WIDTH))


# ============================================================
# 3. Excel出力
# ============================================================

def as_text_cell(cell):
    """文字列のセルを、数式として解釈されない文字セルに固定する。

    openpyxlは「=」で始まる文字列を数式として保存する。申し送りの原文のような外部由来の
    文字列が数式になることを防ぐため、文字列は必ず文字セルとして保存する。
    """

    if isinstance(cell.value, str):
        cell.data_type = "s"
    return cell


# Excelは、これらの文字で始まる文字を数式として実行する(CSVを直接開いたとき)。
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def excel_safe_text(value: object) -> object:
    """文字列がExcelで数式として実行されないよう、危険な文字で始まる場合は先頭に「'」を付ける。"""

    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def excel_safe_csv_path(machine_csv_path: Path) -> Path:
    """機械が読むCSVに対応する、Excelで開いても安全なCSVのパス(例: audit_result_for_excel.csv)。"""

    return machine_csv_path.with_name(f"{machine_csv_path.stem}_for_excel{machine_csv_path.suffix}")


def write_excel_safe_csv(path: Path, records: List[Dict[str, str]]) -> None:
    """Excelで直接開いても数式が実行されないCSVを出力する(人が開く用。採点などの元データには使わない)。"""

    if not records:
        raise ValueError("出力対象のデータがありません")

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=list(records[0].keys()))
        writer.writeheader()
        for record in records:
            writer.writerow({key: excel_safe_text(value) for key, value in record.items()})


def write_audit_xlsx(path: Path, records: List[Dict[str, str]]) -> None:
    """1行目に英語の項目名、2行目に日本語見出しを置いたExcelを出力する。"""

    if not records:
        raise ValueError("出力対象のデータがありません")

    columns = list(records[0].keys())
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "監査結果"

    english_font = Font(color="808080", size=8)
    japanese_font = Font(bold=True)
    japanese_fill = PatternFill("solid", fgColor="E7EEF7")
    # 行の途中で折り返さず1行に収める(長すぎる文だけ折り返す)。上下は中央にそろえる
    cell_alignment = Alignment(wrap_text=True, vertical="center")

    for index, column in enumerate(columns, start=1):
        english_cell = as_text_cell(sheet.cell(row=ENGLISH_ROW, column=index, value=column))
        english_cell.font = english_font
        japanese_cell = as_text_cell(
            sheet.cell(row=JAPANESE_ROW, column=index, value=JAPANESE_LABELS.get(column, column))
        )
        japanese_cell.font = japanese_font
        japanese_cell.fill = japanese_fill
        japanese_cell.alignment = Alignment(vertical="center")

    for row_offset, record in enumerate(records):
        for index, column in enumerate(columns, start=1):
            cell = as_text_cell(
                sheet.cell(
                    row=FIRST_DATA_ROW + row_offset,
                    column=index,
                    value=display_value(column, record.get(column, "")),
                )
            )
            cell.alignment = cell_alignment

    for index, column in enumerate(columns, start=1):
        values = [display_value(column, record.get(column, "")) for record in records]
        sheet.column_dimensions[get_column_letter(index)].width = calculate_column_width(
            column, JAPANESE_LABELS.get(column, column), values
        )
    # 行の高さを内容から決めてそろえる(自動任せだと行ごとにばらつき、余白が目立つため)。
    # 基本は全行を同じ高さ(最大2行分)にし、それ以上必要な行だけ高くする
    needed_lines = []
    for record in records:
        lines = 1
        for index, column in enumerate(columns, start=1):
            width = sheet.column_dimensions[get_column_letter(index)].width
            text_width = display_width(display_value(column, record.get(column, "")))
            lines = max(lines, math.ceil(text_width / max(width - 1, 1)))
        needed_lines.append(lines)
    base_lines = min(TARGET_LINES, max(needed_lines))
    for row_offset, lines in enumerate(needed_lines):
        sheet.row_dimensions[FIRST_DATA_ROW + row_offset].height = (
            ROW_HEIGHT_PER_LINE * max(base_lines, lines)
        )
    sheet.row_dimensions[JAPANESE_ROW].height = 24

    # 見出し2行と識別番号の列を固定し、絞り込み(フィルター)を日本語見出し行へ付ける
    sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=2)
    last_row = FIRST_DATA_ROW + len(records) - 1
    sheet.auto_filter.ref = (
        f"A{JAPANESE_ROW}:{get_column_letter(len(columns))}{last_row}"
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
